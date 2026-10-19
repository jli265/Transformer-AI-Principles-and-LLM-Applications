import warnings
import sys
import os
import threading
import multiprocessing
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext

# ================= 核心配置：完全离线断网模式 =================
# 1. 强行开启本地离线开关（彻底禁止任何联网/在线校验请求）
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"

# 2. 指定本地两个模型的文件夹绝对路径
LLM_MODEL_PATH = r"D:\work\AI Agent\文本大模型2\qwen_model"  # Qwen 大模型路径
EMBED_MODEL_PATH = r"D:\work\AI Agent\文本大模型2\m3e_model"  # m3e 向量模型路径

# 3. 项目本地缓存重定向
PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
D_DRIVE_CACHE_DIR = os.path.join(PROJECT_DIR, "hf_cache")
os.environ["HF_HOME"] = D_DRIVE_CACHE_DIR
# =============================================================

import PyPDF2
from docx import Document
import opencc
import torch
from transformers import AutoTokenizer, AutoModelForCausalLM, pipeline

# 兼容性导入
try:
    from langchain_community.llms import HuggingFacePipeline
    from langchain_community.embeddings import HuggingFaceEmbeddings
    from langchain_community.vectorstores import FAISS
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_core.prompts import ChatPromptTemplate
    from langchain.chains.combine_documents import create_stuff_documents_chain
    from langchain.chains.retrieval import create_retrieval_chain
except ImportError:
    from langchain_community.llms import HuggingFacePipeline
    from langchain_community.embeddings import HuggingFaceEmbeddings
    from langchain_community.vectorstores import FAISS
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_core.prompts import ChatPromptTemplate
    from langchain.chains import create_retrieval_chain
    from langchain.chains.combine_documents import create_stuff_documents_chain


class PureOfflineRAGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("本地离线 AI 论文助手 (纯断网运行版)")
        self.root.geometry("800x600")
        self.chat_history = []
        self.rag_chain = None
        self.cc = opencc.OpenCC('t2s')
        self.setup_ui()
        threading.Thread(target=self.async_load_engine, daemon=True).start()

    def setup_ui(self):
        # 顶部：上传按钮
        top_frame = tk.Frame(self.root)
        top_frame.pack(fill=tk.X, padx=10, pady=10)
        self.btn_upload = tk.Button(top_frame, text="上传 PDF/Word 论文", command=self.upload_file, state=tk.DISABLED,
                                    bg="#4CAF50", fg="white")
        self.btn_upload.pack()

        # 中部：对话框
        self.chat_area = scrolledtext.ScrolledText(self.root, state=tk.DISABLED, font=("Arial", 11))
        self.chat_area.pack(fill=tk.BOTH, expand=True, padx=10)

        # 底部：输入与按钮
        input_frame = tk.Frame(self.root)
        input_frame.pack(fill=tk.X, padx=10, pady=10)
        self.user_input = tk.Text(input_frame, height=3)
        self.user_input.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # 发送与保存按钮区域
        btn_frame = tk.Frame(input_frame)
        btn_frame.pack(side=tk.RIGHT, padx=(10, 0))
        self.btn_send = tk.Button(btn_frame, text="发送", command=self.send_question, state=tk.DISABLED, bg="#2196F3",
                                  fg="white", width=8)
        self.btn_send.pack(pady=2)
        self.btn_save = tk.Button(btn_frame, text="保存为Word", command=self.save_to_word, state=tk.DISABLED, width=8)
        self.btn_save.pack(pady=2)

        self.lbl_status = tk.Label(self.root, text="正在纯本地环境加载算法引擎...", fg="red")
        self.lbl_status.pack(pady=5)

    def async_load_engine(self):
        try:
            device = "cuda" if torch.cuda.is_available() else "cpu"

            # CPU 多核满载计算优化
            if device == "cpu":
                torch.set_num_threads(multiprocessing.cpu_count())

            # 1. 绝对离线加载本地向量模型
            self.embeddings_engine = HuggingFaceEmbeddings(
                model_name=EMBED_MODEL_PATH,
                model_kwargs={'device': device}
            )

            # 2. 绝对离线加载本地 Qwen 大模型
            tokenizer = AutoTokenizer.from_pretrained(
                LLM_MODEL_PATH,
                use_fast=True,
                local_files_only=True
            )
            model = AutoModelForCausalLM.from_pretrained(
                LLM_MODEL_PATH,
                torch_dtype=torch.float16 if device == "cuda" else torch.float32,
                device_map="auto" if device == "cuda" else None,
                low_cpu_mem_usage=True,
                local_files_only=True,
                use_cache=True
            )

            pipe = pipeline(
                "text-generation",
                model=model,
                tokenizer=tokenizer,
                max_new_tokens=300,
                do_sample=False,
                repetition_penalty=1.05,
                pad_token_id=tokenizer.eos_token_id
            )
            self.llm = HuggingFacePipeline(pipeline=pipe)

            status_text = f"引擎已全离线就绪！运行于 {'GPU' if device == 'cuda' else 'CPU'} 模式"
            self.root.after(0, lambda: self.btn_upload.config(state=tk.NORMAL))
            self.root.after(0, lambda: self.lbl_status.config(text=status_text, fg="green"))
        except Exception as e:
            self.root.after(0, lambda err=e: messagebox.showerror(
                "纯离线加载失败",
                f"请检查模型路径是否全部下载就绪：\n1. Qwen路径: {LLM_MODEL_PATH}\n2. Vector路径: {EMBED_MODEL_PATH}\n\n详细错误：{str(err)}"
            ))

    def upload_file(self):
        filepath = filedialog.askopenfilename(filetypes=[("支持的文件", "*.pdf *.docx *.txt")])
        if not filepath: return
        self.lbl_status.config(text="正在解析文档并构建本地向量库...", fg="orange")
        threading.Thread(target=self._process_file_thread, args=(filepath,), daemon=True).start()

    def _process_file_thread(self, filepath):
        try:
            raw_text = ""
            if filepath.endswith('.pdf'):
                with open(filepath, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    raw_text = "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
            elif filepath.endswith('.docx'):
                doc = Document(filepath)
                full_text = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip(): full_text.append(cell.text)
                raw_text = "\n".join(full_text)
            else:
                with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                    raw_text = f.read()

            simp_text = self.cc.convert(raw_text)
            splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
            vectorstore = FAISS.from_texts(splitter.split_text(simp_text), self.embeddings_engine)

            qa_prompt = ChatPromptTemplate.from_template("已知论文内容：\n{context}\n\n问题：{input}\n\n请简洁准确回答：")
            self.rag_chain = create_retrieval_chain(
                vectorstore.as_retriever(search_kwargs={"k": 2}),
                create_stuff_documents_chain(self.llm, qa_prompt)
            )

            self.root.after(0, self._on_file_processed)
        except Exception as e:
            self.root.after(0, lambda err=e: messagebox.showerror("读取失败", str(err)))

    def _on_file_processed(self):
        self.btn_send.config(state=tk.NORMAL)
        self.btn_save.config(state=tk.NORMAL)
        self.lbl_status.config(text="文档构建完成，可以提问", fg="blue")

    def send_question(self):
        q = self.user_input.get(1.0, tk.END).strip()
        if not q: return
        self.user_input.delete(1.0, tk.END)
        self.btn_send.config(state=tk.DISABLED)
        self.lbl_status.config(text="AI 思考推理中...", fg="orange")
        threading.Thread(target=self._run_chain, args=(q,), daemon=True).start()

    def _run_chain(self, q):
        try:
            with torch.inference_mode():
                res = self.rag_chain.invoke({"input": q})
            ans = res['answer']
            self.chat_history.append({"q": q, "a": ans})
            self.root.after(0, self._display_ans, q, ans)
        except Exception as e:
            self.root.after(0, lambda err=e: messagebox.showerror("推理异常", str(err)))
            self.root.after(0, lambda: self.btn_send.config(state=tk.NORMAL))

    def _display_ans(self, q, ans):
        self.chat_area.config(state=tk.NORMAL)
        self.chat_area.insert(tk.END, f"用户: {q}\nAI: {ans}\n\n")
        self.chat_area.yview(tk.END)
        self.chat_area.config(state=tk.DISABLED)
        self.btn_send.config(state=tk.NORMAL)
        self.lbl_status.config(text="回答完毕", fg="blue")

    def save_to_word(self):
        if not self.chat_history: return
        filepath = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word文档", "*.docx")])
        if filepath:
            doc = Document()
            doc.add_heading('AI 论文阅读对话记录', 0)
            for item in self.chat_history:
                doc.add_paragraph(f"用户: {item['q']}", style='Strong')
                doc.add_paragraph(f"AI: {item['a']}")
                doc.add_paragraph("-" * 30)
            doc.save(filepath)
            messagebox.showinfo("成功", "对话已成功保存为 Word 文档！")


if __name__ == "__main__":
    root = tk.Tk()
    app = PureOfflineRAGUI(root)
    root.mainloop()